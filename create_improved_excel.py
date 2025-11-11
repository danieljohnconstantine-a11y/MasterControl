#!/usr/bin/env python3
"""
Create improved Excel spreadsheet with clear 1-row-per-dog format
Focuses on speed metrics and easy comparison
"""

import pandas as pd
import ast
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def extract_speed_metrics(recent_runs_str, distance):
    """Extract speed-related metrics from recent runs"""
    metrics = {
        'BestTime_sec': None,
        'Last3AvgTime_sec': None,
        'BestSectional_sec': None,
        'SpeedRating_kmh': None,
    }
    
    if pd.isna(recent_runs_str):
        return metrics
    
    try:
        runs = ast.literal_eval(recent_runs_str)
    except:
        return metrics
    
    if not runs:
        return metrics
    
    # Extract times at or near this distance (±50m tolerance)
    race_times = []
    sectional_times = []
    
    for run in runs:
        run_dist = run.get('distance')
        if run_dist:
            try:
                run_dist = int(run_dist)
                # Check if distance is close to target
                if abs(run_dist - distance) <= 50:
                    # Extract race time
                    if run.get('racetime'):
                        try:
                            # Parse time like "23.45" or "0:23.45"
                            time_str = run['racetime']
                            if ':' in time_str:
                                parts = time_str.split(':')
                                seconds = float(parts[0]) * 60 + float(parts[1])
                            else:
                                seconds = float(time_str)
                            race_times.append(seconds)
                        except:
                            pass
                    
                    # Extract sectional time
                    if run.get('sectime'):
                        try:
                            sectional_times.append(float(run['sectime']))
                        except:
                            pass
            except:
                pass
    
    # Calculate metrics
    if race_times:
        metrics['BestTime_sec'] = round(min(race_times), 2)
        if len(race_times) >= 3:
            metrics['Last3AvgTime_sec'] = round(sum(race_times[:3]) / 3, 2)
        elif len(race_times) > 0:
            metrics['Last3AvgTime_sec'] = round(sum(race_times) / len(race_times), 2)
        
        # Speed rating in km/h
        metrics['SpeedRating_kmh'] = round((distance / metrics['BestTime_sec']) * 3.6, 2)
    
    if sectional_times:
        metrics['BestSectional_sec'] = round(min(sectional_times), 2)
    
    return metrics


def extract_form_metrics(recent_runs_str, track):
    """Extract form metrics from recent runs"""
    metrics = {
        'Last5_AvgPos': None,
        'Last5_Wins': 0,
        'Last5_Places': 0,
        'FormTrend': 'Unknown',
        'DaysSinceLast': None,
        'Consistency': None,
        'TrackWinPct': None,
    }
    
    if pd.isna(recent_runs_str):
        return metrics
    
    try:
        runs = ast.literal_eval(recent_runs_str)
    except:
        return metrics
    
    if not runs:
        return metrics
    
    # Last 5 races
    last_5 = runs[:5]
    positions = []
    
    for run in last_5:
        if run.get('pos'):
            try:
                pos_num = int(''.join(filter(str.isdigit, run['pos'])))
                positions.append(pos_num)
            except:
                pass
    
    if positions:
        metrics['Last5_AvgPos'] = round(sum(positions) / len(positions), 1)
        metrics['Last5_Wins'] = sum(1 for p in positions if p == 1)
        metrics['Last5_Places'] = sum(1 for p in positions if p <= 3)
        
        # Consistency (lower stdev = more consistent)
        if len(positions) > 1:
            import statistics
            stdev = statistics.stdev(positions)
            # Convert to 0-10 scale (inverted - higher is better)
            metrics['Consistency'] = max(0, round(10 - stdev, 1))
        else:
            metrics['Consistency'] = 5.0
        
        # Trend
        if len(positions) >= 3:
            recent = sum(positions[:2]) / 2
            older = sum(positions[2:]) / len(positions[2:])
            if recent < older - 0.5:
                metrics['FormTrend'] = '↑ Improving'
            elif recent > older + 0.5:
                metrics['FormTrend'] = '↓ Declining'
            else:
                metrics['FormTrend'] = '→ Stable'
    
    # Track-specific performance
    if track:
        track_runs = [r for r in runs if r.get('track') and track.lower() in r['track'].lower()]
        if track_runs:
            track_wins = sum(1 for r in track_runs if r.get('pos') and '1st' in r['pos'])
            metrics['TrackWinPct'] = round((track_wins / len(track_runs)) * 100, 1)
    
    return metrics


def create_simplified_excel(input_csv, output_excel):
    """Create simplified, easy-to-read Excel format"""
    
    # Read data
    df = pd.read_csv(input_csv)
    
    print(f"Processing {len(df)} dogs...")
    
    # Extract speed and form metrics for each dog
    speed_data = []
    form_data = []
    
    for idx, row in df.iterrows():
        distance = row.get('Distance', 0)
        track = row.get('Track', '')
        recent_runs = row.get('RecentRuns')
        
        speed_metrics = extract_speed_metrics(recent_runs, distance)
        form_metrics = extract_form_metrics(recent_runs, track)
        
        speed_data.append(speed_metrics)
        form_data.append(form_metrics)
    
    speed_df = pd.DataFrame(speed_data)
    form_df = pd.DataFrame(form_data)
    
    # Build simplified dataframe with clear column structure
    simple_df = pd.DataFrame()
    
    # GROUP 1: IDENTIFICATION
    simple_df['Track'] = df['Track']
    simple_df['Race'] = df['RaceNumber']
    simple_df['Box'] = df['Box']
    simple_df['Dog'] = df['DogName']
    simple_df['Trainer'] = df['Trainer']
    
    # GROUP 2: SPEED METRICS (PROMINENT)
    simple_df['Best_Time_sec'] = speed_df['BestTime_sec']
    simple_df['Last3_Avg_sec'] = speed_df['Last3AvgTime_sec']
    simple_df['Best_Sectional_sec'] = speed_df['BestSectional_sec']
    simple_df['Speed_kmh'] = speed_df['SpeedRating_kmh']
    
    # Calculate speed rank within each race
    simple_df['Speed_Rank'] = simple_df.groupby('Race')['Speed_kmh'].rank(ascending=False, method='min')
    
    # GROUP 3: RECENT FORM
    simple_df['L5_Avg_Pos'] = form_df['Last5_AvgPos']
    simple_df['L5_Wins'] = form_df['Last5_Wins']
    simple_df['L5_Places'] = form_df['Last5_Places']
    simple_df['Form_Trend'] = form_df['FormTrend']
    simple_df['Consistency'] = form_df['Consistency']
    simple_df['Track_Win%'] = form_df['TrackWinPct']
    
    # GROUP 4: CAREER STATS
    simple_df['Career_Wins'] = df['CareerWins']
    simple_df['Career_Starts'] = df['CareerStarts']
    simple_df['Win%'] = df.get('WinPercent', None)
    simple_df['Prize_Money'] = df['PrizeMoney']
    simple_df['API'] = df.get('API', None)
    
    # GROUP 5: BETTING SCORE
    simple_df['Final_Score'] = df['FinalScore']
    
    # Add confidence based on data completeness
    def calc_confidence(row):
        score = 0
        if pd.notna(row['Best_Time_sec']): score += 1
        if pd.notna(row['L5_Avg_Pos']): score += 1
        if pd.notna(row['Speed_kmh']): score += 1
        if row['Career_Starts'] > 10: score += 1
        
        if score >= 3:
            return 'High'
        elif score >= 2:
            return 'Medium'
        else:
            return 'Low'
    
    simple_df['Confidence'] = simple_df.apply(calc_confidence, axis=1)
    
    # Create Excel with multiple views
    with pd.ExcelWriter(output_excel, engine='openpyxl') as writer:
        
        # Sheet 1: Race View (sorted by race and box)
        race_view = simple_df.sort_values(['Race', 'Box'])
        race_view.to_excel(writer, sheet_name='Race View', index=False)
        
        # Sheet 2: Speed Rankings (sorted by speed within race)
        speed_view = simple_df.sort_values(['Race', 'Speed_kmh'], ascending=[True, False])
        speed_view.to_excel(writer, sheet_name='Speed Rankings', index=False)
        
        # Sheet 3: Form Rankings (sorted by form)
        form_view = simple_df.sort_values(['Race', 'L5_Avg_Pos'], ascending=[True, True])
        form_view.to_excel(writer, sheet_name='Form Rankings', index=False)
        
        # Sheet 4: Top Picks (best per race)
        picks = simple_df.sort_values(['Race', 'Final_Score'], ascending=[True, False])
        picks = picks.groupby('Race').first().reset_index()
        def format_pick_reason(r):
            speed = f"{r['Speed_kmh']:.1f}km/h" if pd.notna(r['Speed_kmh']) else 'N/A'
            rank = f"Rank {int(r['Speed_Rank'])}" if pd.notna(r['Speed_Rank']) else 'Rank N/A'
            form = f"{r['L5_Avg_Pos']:.1f} avg" if pd.notna(r['L5_Avg_Pos']) else 'N/A'
            trend = r['Form_Trend'] if pd.notna(r['Form_Trend']) else 'Unknown'
            return f"Speed: {speed} ({rank}), Form: {form}, Trend: {trend}"
        
        picks['Pick_Reason'] = picks.apply(format_pick_reason, axis=1)
        picks.to_excel(writer, sheet_name='Top Picks', index=False)
    
    # Apply formatting
    wb = load_workbook(output_excel)
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        # Header formatting
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=11)
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Freeze panes (freeze first row and first 4 columns)
        ws.freeze_panes = 'E2'
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Add borders between race groups (for Race View sheet)
        if sheet_name == 'Race View':
            thick_border = Border(
                bottom=Side(style='medium', color='000000')
            )
            
            prev_race = None
            for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                race_val = row[1].value  # Race column
                if prev_race is not None and race_val != prev_race:
                    # Add border to previous row
                    for cell in ws[row_idx - 1]:
                        cell.border = thick_border
                prev_race = race_val
    
    wb.save(output_excel)
    
    print(f"✅ Created: {output_excel}")
    print(f"   - {len(simple_df)} dogs")
    print(f"   - {len(simple_df.columns)} columns")
    print(f"   - 4 sheets (Race View, Speed Rankings, Form Rankings, Top Picks)")
    
    return simple_df


if __name__ == '__main__':
    # Create improved Excel
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    output_file = f'outputs/greyhound_simplified_{timestamp}.xlsx'
    
    df = create_simplified_excel('outputs/todays_form.csv', output_file)
    
    # Show preview
    print("\n" + "=" * 80)
    print("PREVIEW - Race 1 Dogs")
    print("=" * 80)
    race1 = df[df['Race'] == 1].copy()
    race1 = race1.sort_values('Speed_Rank')
    
    preview_cols = ['Box', 'Dog', 'Speed_kmh', 'Speed_Rank', 'Best_Time_sec', 
                    'L5_Avg_Pos', 'Form_Trend', 'Final_Score']
    print(race1[preview_cols].to_string(index=False))
    
    print(f"\n📁 File saved: {output_file}")
