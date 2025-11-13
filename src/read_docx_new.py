"""
Read DOCX files and extract all content as structured blocks.
No parsing - only extraction.
"""
from docx import Document
from typing import List, Dict, Any
import os


def extract_docx_content(docx_path: str) -> List[Dict[str, Any]]:
    """
    Extract all content from a DOCX file as structured blocks.
    
    Returns a list of blocks, each with:
    - type: "paragraph" | "table"
    - text: list of text strings (for paragraphs)
    - table: list of lists (for tables)
    - raw: original text
    """
    blocks = []
    doc = Document(docx_path)
    
    # Extract paragraphs (simple approach - just iterate paragraphs)
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            blocks.append({
                "type": "paragraph",
                "text": [text],
                "raw": text
            })
    
    # Extract tables
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            table_data.append(row_data)
        
        if table_data:
            blocks.append({
                "type": "table",
                "table": table_data,
                "raw": "\n".join(["\t".join(row) for row in table_data])
            })
    
    return blocks
    
    return blocks


def read_all_docx_files(data_dir: str = "data") -> Dict[str, List[Dict[str, Any]]]:
    """
    Read all DOCX files from the data directory.
    
    Returns: dict mapping filename -> list of blocks
    """
    result = {}
    
    if not os.path.exists(data_dir):
        return result
    
    for filename in os.listdir(data_dir):
        if filename.endswith('.docx') and not filename.startswith('~'):
            filepath = os.path.join(data_dir, filename)
            try:
                blocks = extract_docx_content(filepath)
                result[filename] = blocks
            except Exception as e:
                print(f"Error reading {filename}: {e}")
    
    return result
