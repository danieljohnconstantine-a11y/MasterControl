import pandas as pd
import os

def export_to_excel(dogs, output_path):
    """Legacy export function - kept for compatibility"""
    # Flatten list fields
    for dog in dogs:
        dog["recent_positions"] = ", ".join(map(str, dog.get("recent_positions", [])))
        dog["form_trend"] = str(dog.get("form_trend", ""))
        dog["has_win"] = int(dog.get("has_win", 0))
        dog["has_place"] = int(dog.get("has_place", 0))

    # Define strict column order
    columns = [
        "Track", "RaceNumber", "RaceDate", "RaceTime", "Distance",
        "Box", "DogsName", "form_code", "age_sex", "weight", "trainer",
        "wins", "places", "starts", "PrizeMoney", "KmH", "experience_level",
        "FinalScore", "Bet", "strike_rate", "win_percentage", "place_percentage",
        "consistency_rate", "consistent_places", "has_dnf", "has_win", "has_place",
        "recent_races", "recent_positions", "avg_recent_position",
        "best_recent_position", "worst_recent_position", "form_trend",
        "source_file", "Date"
    ]

    # Fill missing keys with None
    for dog in dogs:
        for col in columns:
            if col not in dog:
                dog[col] = None

    df = pd.DataFrame(dogs)[columns]
    filename = f"greyhound_analysis_{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(output_path, filename)
    df.to_excel(filepath, index=False)
    print(f"[OK] EXCEL SAVED: {filepath}")


def export_ordered(df, output_dir):
    """
    Export ordered and sorted DataFrame to both CSV and XLSX with formatting.
    
    Args:
        df: DataFrame to export
        output_dir: Directory to save files
    
    Returns:
        tuple: (csv_path, xlsx_path)
    """
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    
    # PRE-EXPORT VALIDATION: Check for missing critical columns
    critical_cols = ['S2_1_Distance', 'S2_1_RaceTime', 'Speed_kmh']
    missing_data_warning = False
    
    for col in critical_cols:
        if col in df.columns:
            nan_pct = (df[col].isna().sum() / len(df)) * 100
            if nan_pct > 10:
                print(f"[WARN] Column '{col}' has {nan_pct:.1f}% missing values (>10% threshold)")
                missing_data_warning = True
    
    if missing_data_warning:
        print("[WARN] Incomplete Section 2 detected. Consider reviewing parser logic.")
        print("[INFO] Proceeding with export, but data quality may be compromised.")
    
    # 1) Reorder columns
    priority_cols = ["Track", "RaceNumber", "Box", "DogName", "FinalScore", "Speed_kmh"]
    # Only include priority cols that exist
    priority_cols = [c for c in priority_cols if c in df.columns]
    remaining = [c for c in df.columns if c not in priority_cols]
    df = df[priority_cols + remaining]

    # 2) Sort by Track -> RaceNumber -> Box
    sort_cols = []
    if "Track" in df.columns:
        sort_cols.append("Track")
    if "RaceNumber" in df.columns:
        sort_cols.append("RaceNumber")
    if "Box" in df.columns:
        sort_cols.append("Box")
    
    if sort_cols:
        df = df.sort_values(by=sort_cols, ascending=[True] * len(sort_cols))

    # 3) Save files
    os.makedirs(output_dir, exist_ok=True)
    csv_out = os.path.join(output_dir, "greyhound_comparison_ordered.csv")
    xlsx_out = os.path.join(output_dir, "greyhound_comparison_ordered.xlsx")
    df.to_csv(csv_out, index=False)
    df.to_excel(xlsx_out, index=False, engine='openpyxl')

    # 4) Excel formatting
    try:
        wb = load_workbook(xlsx_out)
        ws = wb.active
        ws.title = "Greyhounds"
        ws.freeze_panes = "A2"
        
        for i, col_name in enumerate(df.columns, 1):
            width = min(40, max(10, len(str(col_name)) + 2))
            ws.column_dimensions[get_column_letter(i)].width = width
            
            # Apply number format to numeric columns
            if col_name in df.columns and df[col_name].dtype.kind in "fi":
                for row_idx in range(2, len(df) + 2):  # Skip header
                    cell = ws[f"{get_column_letter(i)}{row_idx}"]
                    cell.number_format = "0.00"
        
        wb.save(xlsx_out)
    except Exception as e:
        print(f"[WARN] Excel formatting failed: {e}")

    print(f"[OK] Exported {len(df)} rows x {len(df.columns)} columns (ordered by Track/Race/Box)")
    print(f"[FILES] {csv_out} | {xlsx_out}")

    return csv_out, xlsx_out
